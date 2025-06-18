#!/usr/bin/env python3
"""
Unit tests for Excel formula conversion to Python functions.

This test suite:
1. Extracts all data from the original Excel file to CSV
2. Runs the script generated from the LLM prompt
3. Compares calculated columns with expected results
"""

import pytest
import pandas as pd
import openpyxl
import subprocess
import sys
from pathlib import Path
import tempfile
import shutil
import os


class TestFormulaConversion:
    """Test class for Excel formula conversion validation."""
    
    @pytest.fixture(scope="class")
    def setup_test_environment(self):
        """Set up test environment with sample data and expected results."""
        # Create temporary directory for test files
        self.temp_dir = Path(tempfile.mkdtemp())
        self.excel_file = self.temp_dir / "test_data.xlsx"
        self.input_csv = self.temp_dir / "input_data.csv"
        self.expected_csv = self.temp_dir / "expected_data.csv"
        self.output_csv = self.temp_dir / "output_data.csv"
        self.generated_script = self.temp_dir / "generated_script.py"
        
        # Create sample Excel file with formulas
        self._create_sample_excel()
        
        # Extract all data including calculated values
        self._extract_all_data_to_csv()
        
        # Extract input data (non-formula columns only)
        self._extract_input_data_to_csv()
        
        yield {
            'temp_dir': self.temp_dir,
            'excel_file': self.excel_file,
            'input_csv': self.input_csv,
            'expected_csv': self.expected_csv,
            'output_csv': self.output_csv,
            'generated_script': self.generated_script
        }
        
        # Cleanup
        shutil.rmtree(self.temp_dir)
    
    def _create_sample_excel(self):
        """Create sample Excel file with formulas for testing."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Test Data"
        
        # Headers
        headers = ['Name', 'Quantity', 'Price', 'Total', 'Discount%', 'Final_Price', 'Tax_Rate', 'With_Tax']
        for i, header in enumerate(headers, 1):
            ws.cell(row=1, column=i, value=header)
        
        # Sample data
        sample_data = [
            ['Product A', 10, 25.50],
            ['Product B', 5, 15.75],  
            ['Product C', 8, 30.00],
            ['Product D', 12, 8.25],
            ['Product E', 6, 45.00]
        ]
        
        for row_idx, row_data in enumerate(sample_data, 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # Add formulas
        for row in range(2, 7):
            # Total = Quantity * Price (Column D)
            ws.cell(row=row, column=4, value=f'=B{row}*C{row}')
            
            # Discount percentage (Column E - fixed 10%)
            ws.cell(row=row, column=5, value=0.1)
            
            # Final_Price = Total * (1 - Discount%) (Column F)
            ws.cell(row=row, column=6, value=f'=D{row}*(1-E{row})')
            
            # Tax rate (Column G - fixed 8.5%)
            ws.cell(row=row, column=7, value=0.085)
            
            # With_Tax = Final_Price * (1 + Tax_Rate) (Column H)
            ws.cell(row=row, column=8, value=f'=F{row}*(1+G{row})')
        
        wb.save(self.excel_file)
    
    def _extract_all_data_to_csv(self):
        """Extract all data including calculated values to CSV for comparison."""
        # Load with data_only=False to get formulas first
        wb = openpyxl.load_workbook(self.excel_file, data_only=False)
        ws = wb.active
        
        # Get all data
        data = []
        headers = []
        
        # Get headers
        for col_idx in range(1, ws.max_column + 1):
            header_cell = ws.cell(row=1, column=col_idx)
            headers.append(header_cell.value)
        
        # Get data rows and calculate values programmatically
        for row_idx in range(2, ws.max_row + 1):
            row_data = []
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                
                # If it's a formula cell, calculate it manually
                if cell.data_type == 'f':
                    formula = cell.value
                    if formula == f'=B{row_idx}*C{row_idx}':
                        # Total = Quantity * Price
                        quantity = ws.cell(row=row_idx, column=2).value
                        price = ws.cell(row=row_idx, column=3).value
                        calculated_value = quantity * price
                        row_data.append(calculated_value)
                    elif formula == f'=D{row_idx}*(1-E{row_idx})':
                        # Final_Price = Total * (1 - Discount%)
                        total = ws.cell(row=row_idx, column=4).value
                        discount = ws.cell(row=row_idx, column=5).value
                        # If total is a formula, calculate it first
                        if ws.cell(row=row_idx, column=4).data_type == 'f':
                            quantity = ws.cell(row=row_idx, column=2).value
                            price = ws.cell(row=row_idx, column=3).value
                            total = quantity * price
                        calculated_value = total * (1 - discount)
                        row_data.append(calculated_value)
                    elif formula == f'=F{row_idx}*(1+G{row_idx})':
                        # With_Tax = Final_Price * (1 + Tax_Rate)
                        final_price = ws.cell(row=row_idx, column=6).value
                        tax_rate = ws.cell(row=row_idx, column=7).value
                        # If final_price is a formula, calculate it first
                        if ws.cell(row=row_idx, column=6).data_type == 'f':
                            quantity = ws.cell(row=row_idx, column=2).value
                            price = ws.cell(row=row_idx, column=3).value
                            discount = ws.cell(row=row_idx, column=5).value
                            total = quantity * price
                            final_price = total * (1 - discount)
                        calculated_value = final_price * (1 + tax_rate)
                        row_data.append(calculated_value)
                    else:
                        # Unknown formula - try to get computed value
                        wb_calc = openpyxl.load_workbook(self.excel_file, data_only=True)
                        ws_calc = wb_calc.active
                        calculated_value = ws_calc.cell(row=row_idx, column=col_idx).value
                        row_data.append(calculated_value)
                else:
                    # Regular cell
                    row_data.append(cell.value)
            data.append(row_data)
        
        # Create DataFrame and save
        df = pd.DataFrame(data, columns=headers)
        df.to_csv(self.expected_csv, index=False)
    
    def _extract_input_data_to_csv(self):
        """Extract only non-formula columns to CSV for input to generated script."""
        # Load without data_only to check for formulas
        wb = openpyxl.load_workbook(self.excel_file, data_only=False)
        ws = wb.active
        
        # Identify non-formula columns
        input_columns = []
        headers = []
        
        # Get headers and identify non-formula columns
        for col_idx in range(1, ws.max_column + 1):
            header_cell = ws.cell(row=1, column=col_idx)
            header = header_cell.value
            headers.append(header)
            
            # Check if this column has formulas
            has_formula = False
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.data_type == 'f':  # Formula cell
                    has_formula = True
                    break
            
            if not has_formula:
                input_columns.append(col_idx)
        
        # Extract data for non-formula columns
        data = []
        input_headers = [headers[col_idx-1] for col_idx in input_columns]
        
        for row_idx in range(2, ws.max_row + 1):
            row_data = []
            for col_idx in input_columns:
                cell = ws.cell(row=row_idx, column=col_idx)
                row_data.append(cell.value)
            data.append(row_data)
        
        # Create DataFrame and save
        df = pd.DataFrame(data, columns=input_headers)
        df.to_csv(self.input_csv, index=False)
    
    def test_prompt_generation(self, setup_test_environment):
        """Test that the prompt generator works correctly."""
        env = setup_test_environment
        
        # Run prompt generator
        from src.prompt_generator import ExcelFormulaPromptGenerator
        
        generator = ExcelFormulaPromptGenerator(env['excel_file'])
        prompt = generator.process(env['temp_dir'] / 'test_prompt.txt')
        
        # Verify prompt contains expected elements
        assert "Excel Formula to Python Function Conversion Request" in prompt
        assert "Data Schema and Sample" in prompt
        assert "Formulas to Convert" in prompt
        assert "Required Output" in prompt
        
        # Check that formulas were found
        assert len(generator.unique_formulas) > 0
        print(f"Found {len(generator.unique_formulas)} unique formulas")
    
    def test_sample_generated_script(self, setup_test_environment):
        """Test with a sample generated script to verify the test framework works."""
        env = setup_test_environment
        
        # Create a sample script that mimics what the LLM would generate
        sample_script = f'''
import pandas as pd

def calculate_total(row):
    """Calculate Total = Quantity * Price"""
    return row['Quantity'] * row['Price']

def calculate_final_price(row):
    """Calculate Final_Price = Total * (1 - Discount%)"""
    return row['Total'] * (1 - row['Discount%'])

def calculate_with_tax(row):
    """Calculate With_Tax = Final_Price * (1 + Tax_Rate)"""
    return row['Final_Price'] * (1 + row['Tax_Rate'])

# Main execution
if __name__ == "__main__":
    # Read input data
    df = pd.read_csv('{env['input_csv']}')
    
    # Apply formulas to create calculated columns
    df['Total'] = df.apply(calculate_total, axis=1)
    df['Final_Price'] = df.apply(calculate_final_price, axis=1)
    df['With_Tax'] = df.apply(calculate_with_tax, axis=1)
    
    # Save result
    df.to_csv('{env['output_csv']}', index=False)
    print("Script executed successfully!")
'''
        
        # Save the sample script
        with open(env['generated_script'], 'w') as f:
            f.write(sample_script)
        
        # Run the generated script
        result = subprocess.run([sys.executable, str(env['generated_script'])], 
                              capture_output=True, text=True, cwd=env['temp_dir'])
        
        assert result.returncode == 0, f"Script failed with error: {result.stderr}"
        assert env['output_csv'].exists(), "Output CSV file was not created"
        
        # Load expected and actual results
        expected_df = pd.read_csv(env['expected_csv'])
        actual_df = pd.read_csv(env['output_csv'])
        
        # Compare calculated columns
        formula_columns = ['Total', 'Final_Price', 'With_Tax']
        
        for col in formula_columns:
            if col in expected_df.columns and col in actual_df.columns:
                # Compare with tolerance for floating point precision
                pd.testing.assert_series_equal(
                    expected_df[col], 
                    actual_df[col], 
                    check_names=True,
                    rtol=1e-10,  # Relative tolerance
                    atol=1e-10   # Absolute tolerance  
                )
                print(f"✓ Column '{col}' matches expected values")
    
    def test_excel_data_extraction(self, setup_test_environment):
        """Test that Excel data extraction works correctly."""
        env = setup_test_environment
        
        # Verify input CSV was created and has expected structure
        assert env['input_csv'].exists(), "Input CSV file was not created"
        
        input_df = pd.read_csv(env['input_csv'])
        
        # Should contain non-formula columns
        expected_input_columns = ['Name', 'Quantity', 'Price', 'Discount%', 'Tax_Rate']
        for col in expected_input_columns:
            assert col in input_df.columns, f"Missing expected column: {col}"
        
        # Should not contain formula columns
        formula_columns = ['Total', 'Final_Price', 'With_Tax']
        for col in formula_columns:
            assert col not in input_df.columns, f"Input CSV should not contain calculated column: {col}"
        
        # Verify expected CSV contains all columns
        expected_df = pd.read_csv(env['expected_csv'])
        all_expected_columns = expected_input_columns + formula_columns
        
        for col in all_expected_columns:
            assert col in expected_df.columns, f"Missing expected column in full data: {col}"
        
        print("✓ Excel data extraction working correctly")
    
    def test_formula_identification(self, setup_test_environment):
        """Test that formulas are correctly identified and extracted."""
        env = setup_test_environment
        
        from src.prompt_generator import ExcelFormulaPromptGenerator
        
        generator = ExcelFormulaPromptGenerator(env['excel_file'])
        generator.load_workbook()
        generator.create_header_mapping()
        generator.extract_formulas()
        
        # Should find 3 unique formulas
        expected_formulas = [
            '=B2*C2',  # Total calculation (will be same pattern for all rows)
            '=D2*(1-E2)',  # Final price calculation  
            '=F2*(1+G2)'   # With tax calculation
        ]
        
        # Convert to set for comparison (ignoring row numbers)
        found_formula_patterns = set()
        for formula in generator.unique_formulas:
            # Normalize row numbers to 2 for comparison
            normalized = formula.replace('3', '2').replace('4', '2').replace('5', '2').replace('6', '2')
            found_formula_patterns.add(normalized)
        
        for expected in expected_formulas:
            assert expected in found_formula_patterns, f"Expected formula pattern not found: {expected}"
        
        print(f"✓ Found expected formula patterns: {found_formula_patterns}")


# Additional utility function for manual testing
def run_full_test_cycle(excel_file_path, generated_script_path):
    """
    Run a full test cycle with a real generated script.
    
    Args:
        excel_file_path: Path to Excel file with formulas
        generated_script_path: Path to Python script generated by LLM
    """
    temp_dir = Path(tempfile.mkdtemp())
    
    try:
        # Setup paths
        input_csv = temp_dir / "input_data.csv"
        expected_csv = temp_dir / "expected_data.csv"
        output_csv = temp_dir / "output_data.csv"
        
        # Extract data from Excel
        test_instance = TestFormulaConversion()
        test_instance.temp_dir = temp_dir
        test_instance.excel_file = Path(excel_file_path)
        test_instance.input_csv = input_csv
        test_instance.expected_csv = expected_csv
        test_instance.output_csv = output_csv
        
        test_instance._extract_all_data_to_csv()
        test_instance._extract_input_data_to_csv()
        
        # Run the generated script
        result = subprocess.run([sys.executable, generated_script_path], 
                              capture_output=True, text=True, cwd=temp_dir)
        
        if result.returncode != 0:
            print(f"Script failed with error: {result.stderr}")
            return False
        
        # Compare results
        expected_df = pd.read_csv(expected_csv)
        actual_df = pd.read_csv(output_csv)
        
        print("Comparison Results:")
        print("==================")
        
        # Find calculated columns by comparing with input
        input_df = pd.read_csv(input_csv)
        calculated_columns = [col for col in expected_df.columns if col not in input_df.columns]
        
        success = True
        for col in calculated_columns:
            if col in actual_df.columns:
                try:
                    pd.testing.assert_series_equal(
                        expected_df[col], 
                        actual_df[col], 
                        check_names=True,
                        rtol=1e-10,
                        atol=1e-10
                    )
                    print(f"✓ Column '{col}': PASS")
                except AssertionError as e:
                    print(f"✗ Column '{col}': FAIL - {str(e)}")
                    success = False
            else:
                print(f"✗ Column '{col}': MISSING from output")
                success = False
        
        return success
        
    finally:
        shutil.rmtree(temp_dir)


if __name__ == "__main__":
    # Run tests
    pytest.main([__file__, "-v"])
