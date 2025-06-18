import openpyxl
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Create sample Excel file with formulas
wb = Workbook()
ws = wb.active
ws.title = "Sample Data"

# Add headers
headers = ['Name', 'Quantity', 'Price', 'Total', 'Discount%', 'Final_Price', 'Tax_Rate', 'With_Tax']
for i, header in enumerate(headers, 1):
    ws.cell(row=1, column=i, value=header)

# Add sample data
sample_data = [
    ['Product A', 10, 25.50],
    ['Product B', 5, 15.75],  
    ['Product C', 8, 30.00],
    ['Product D', 12, 8.25],
    ['Product E', 6, 45.00]
]

for row_idx, row_data in enumerate(sample_data, 2):
    for col_idx, value in enumerate(row_data, 1):
        ws.cell(row=row_idx, column=col_idx, value=value)

# Add formulas (columns D, F, H have formulas)
for row in range(2, 7):  # rows 2-6
    # Total = Quantity * Price
    ws.cell(row=row, column=4, value=f'=B{row}*C{row}')
    
    # Discount percentage (fixed 10%)
    ws.cell(row=row, column=5, value=0.1)
    
    # Final_Price = Total * (1 - Discount%)
    ws.cell(row=row, column=6, value=f'=D{row}*(1-E{row})')
    
    # Tax rate (fixed 8.5%)
    ws.cell(row=row, column=7, value=0.085)
    
    # With_Tax = Final_Price * (1 + Tax_Rate)
    ws.cell(row=row, column=8, value=f'=F{row}*(1+G{row})')

wb.save('sample_data.xlsx')
print("Sample Excel file created successfully!") 