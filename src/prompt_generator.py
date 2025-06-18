#!/usr/bin/env python3
"""
Excel Formula to Python Function Prompt Generator

This script analyzes Excel files containing formulas and generates prompts for LLMs
to convert those formulas into Python functions that can be applied to pandas DataFrames.
"""

import openpyxl
import pandas as pd
import re
from openpyxl.utils import get_column_letter, column_index_from_string
from pathlib import Path
import argparse
import sys


class ExcelFormulaPromptGenerator:
    def __init__(self, excel_file_path):
        self.excel_file_path = Path(excel_file_path)
        self.workbook = None
        self.worksheet = None
        self.headers = {}
        self.column_to_header = {}
        self.header_to_column = {}
        self.formulas = {}
        self.unique_formulas = set()
        self.data_columns = {}
        self.formula_patterns = {}  # Pattern -> column info
        
    def load_workbook(self):
        """Load the Excel workbook and get the active worksheet."""
        try:
            self.workbook = openpyxl.load_workbook(self.excel_file_path, data_only=False)
            self.worksheet = self.workbook.active
            print(f"Loaded workbook: {self.excel_file_path}")
            print(f"Active worksheet: {self.worksheet.title}")
        except Exception as e:
            print(f"Error loading workbook: {e}")
            sys.exit(1)
    
    def create_header_mapping(self):
        """Create mapping between headers and Excel column names (A, B, C, etc.)."""
        # Get headers from first row
        for col_idx in range(1, self.worksheet.max_column + 1):
            col_letter = get_column_letter(col_idx)
            header_cell = self.worksheet.cell(row=1, column=col_idx)
            header_value = header_cell.value
            
            if header_value:
                self.headers[col_idx] = header_value
                self.column_to_header[col_letter] = header_value
                self.header_to_column[header_value] = col_letter
        
        print(f"Header mapping created: {self.column_to_header}")
    
    def extract_formulas(self):
        """Extract all formulas from the worksheet and identify unique patterns."""
        formula_columns = []
        formula_patterns = {}  # Pattern -> column info
        
        for col_idx in range(1, self.worksheet.max_column + 1):
            col_letter = get_column_letter(col_idx)
            has_formula = False
            
            # Check if any cell in this column has a formula
            for row_idx in range(2, self.worksheet.max_row + 1):
                cell = self.worksheet.cell(row=row_idx, column=col_idx)
                if cell.data_type == 'f':  # Formula cell
                    has_formula = True
                    formula = cell.value
                    
                    # Normalize formula by replacing row numbers with generic pattern
                    normalized_formula = self._normalize_formula_pattern(formula)
                    
                    # Store the original formula for this pattern
                    if normalized_formula not in formula_patterns:
                        formula_patterns[normalized_formula] = {
                            'example_formula': formula,
                            'column': col_letter,
                            'column_name': self.headers.get(col_idx, f'Column_{col_letter}')
                        }
                    
                    # Store in our tracking structures
                    if formula not in self.formulas:
                        self.formulas[formula] = []
                    self.formulas[formula].append((row_idx, col_idx))
                    break  # Only need one example per column
            
            if has_formula:
                formula_columns.append(col_letter)
        
        # Store unique patterns instead of all individual formulas
        self.unique_formulas = set(pattern_info['example_formula'] for pattern_info in formula_patterns.values())
        self.formula_patterns = formula_patterns
        
        print(f"Found formulas in columns: {formula_columns}")
        print(f"Unique formula patterns found: {len(self.unique_formulas)}")
        
    def _normalize_formula_pattern(self, formula):
        """Normalize formula by replacing specific row/column references with patterns."""
        # Replace cell references like B2, C3, etc. with generic pattern
        pattern = re.sub(r'([A-Z]+)(\d+)', r'\1{row}', formula)
        return pattern
    
    def replace_cell_references_with_headers(self, formula):
        """Replace cell references in formulas with header names."""
        def replace_cell_ref(match):
            cell_ref = match.group(0)
            col_letter = ''.join(filter(str.isalpha, cell_ref))
            
            if col_letter in self.column_to_header:
                return self.column_to_header[col_letter]
            return cell_ref
        
        # Pattern to match cell references like A1, B2, etc.
        cell_pattern = r'[A-Z]+\d+'
        readable_formula = re.sub(cell_pattern, replace_cell_ref, formula)
        return readable_formula
    
    def extract_data_columns(self):
        """Extract data from columns that don't contain formulas."""
        data_columns = {}
        
        for col_idx in range(1, self.worksheet.max_column + 1):
            col_letter = get_column_letter(col_idx)
            header = self.headers.get(col_idx)
            
            if not header:
                continue
                
            # Check if this column has any formulas
            has_formula = False
            for row_idx in range(2, self.worksheet.max_row + 1):
                cell = self.worksheet.cell(row=row_idx, column=col_idx)
                if cell.data_type == 'f':
                    has_formula = True
                    break
            
            # If no formulas, extract the data
            if not has_formula:
                column_data = []
                for row_idx in range(2, self.worksheet.max_row + 1):
                    cell = self.worksheet.cell(row=row_idx, column=col_idx)
                    column_data.append(cell.value)
                data_columns[header] = column_data
        
        self.data_columns = data_columns
        print(f"Data columns extracted: {list(data_columns.keys())}")
    
    def generate_csv_sample(self, num_rows=10):
        """Generate CSV sample of the data for the prompt."""
        if not self.data_columns:
            return ""
        
        # Create DataFrame from data columns
        df = pd.DataFrame(self.data_columns)
        
        # Limit to specified number of rows
        sample_df = df.head(num_rows)
        
        # Convert to CSV string
        csv_string = sample_df.to_csv(index=False)
        return csv_string
    
    def generate_prompt(self):
        """Generate the complete LLM prompt."""
        # Get unique formulas with readable names and organize by purpose
        organized_formulas = []
        
        for pattern, pattern_info in self.formula_patterns.items():
            formula = pattern_info['example_formula']
            readable = self.replace_cell_references_with_headers(formula)
            
            organized_formulas.append({
                'original': formula,
                'readable': readable,
                'column_name': pattern_info['column_name'],
                'column_letter': pattern_info['column']
            })
        
        # Generate CSV sample
        csv_sample = self.generate_csv_sample(10)
        
        # Build the prompt
        prompt = f"""# Excel Formula to Python Function Conversion Request

## Objective
Convert Excel formulas to Python functions that can be applied to pandas DataFrames using the `pd.DataFrame.apply()` method.

## Data Schema and Sample
The input CSV file contains the following columns with sample data (first 10 rows):

```csv
{csv_sample.strip()}
```

## Formulas to Convert
The following formulas need to be converted to Python functions:

"""
        
        for i, formula_info in enumerate(organized_formulas, 1):
            prompt += f"""### Formula {i}: Calculate {formula_info['column_name']}
- **Target Column**: `{formula_info['column_name']}`
- **Original Excel Formula**: `{formula_info['original']}`
- **With Column Names**: `{formula_info['readable']}`

"""
        
        prompt += f"""## Required Output
Please generate a Python script that:

1. **Reads the CSV file** into a pandas DataFrame with the schema shown above
2. **Defines Python functions** for each of the {len(organized_formulas)} formulas listed above
3. **Creates new columns** using `pd.DataFrame.apply()` to apply these functions
4. **Saves the result** as a new CSV file

### Script Requirements:
- Use `pandas` for data manipulation
- Each formula should be implemented as a separate Python function
- Use `df.apply()` with `axis=1` to apply functions row-wise
- Include proper error handling
- Add comments explaining each function
- The output CSV should contain both original and calculated columns

### Example Structure:
```python
import pandas as pd

def calculate_{organized_formulas[0]['column_name'].lower().replace('%', 'percent').replace(' ', '_')}(row):
    '''Calculate {organized_formulas[0]['column_name']}: {organized_formulas[0]['readable']}'''
    return {self._convert_formula_to_python(organized_formulas[0]['readable'])}

def calculate_{organized_formulas[1]['column_name'].lower().replace('%', 'percent').replace(' ', '_') if len(organized_formulas) > 1 else 'second_formula'}(row):
    '''Calculate {organized_formulas[1]['column_name'] if len(organized_formulas) > 1 else 'Second Formula'}: {organized_formulas[1]['readable'] if len(organized_formulas) > 1 else 'etc.'}'''
    return # your calculation here

# Main execution
if __name__ == "__main__":
    # Read input data
    df = pd.read_csv('input_data.csv')
    
    # Apply formulas to create calculated columns
    df['{organized_formulas[0]['column_name']}'] = df.apply(calculate_{organized_formulas[0]['column_name'].lower().replace('%', 'percent').replace(' ', '_')}, axis=1)
    df['{organized_formulas[1]['column_name'] if len(organized_formulas) > 1 else 'Second_Column'}'] = df.apply(calculate_{organized_formulas[1]['column_name'].lower().replace('%', 'percent').replace(' ', '_') if len(organized_formulas) > 1 else 'second_formula'}, axis=1)
    
    # Save result
    df.to_csv('output_data.csv', index=False)
    print("Conversion completed successfully!")
```

Please provide the complete, runnable Python script that implements all the formulas shown above.
"""
        
        return prompt
    
    def _convert_formula_to_python(self, formula):
        """Convert simple Excel formula to Python code hint."""
        # Remove the = sign
        python_expr = formula.replace('=', '')
        
        # Replace Excel operators with Python equivalents
        python_expr = python_expr.replace('*', ' * ')
        python_expr = python_expr.replace('+', ' + ')
        python_expr = python_expr.replace('-', ' - ')
        python_expr = python_expr.replace('/', ' / ')
        
        # Add row[] references for column names
        import re
        # Find column names (assume they are words)
        words = re.findall(r'[A-Za-z_][A-Za-z0-9_%]*', python_expr)
        for word in words:
            if word in self.column_to_header.values():
                python_expr = python_expr.replace(word, f"row['{word}']")
        
        return python_expr
    
    def save_prompt_to_file(self, prompt, output_file="generated_prompt.txt"):
        """Save the generated prompt to a file."""
        with open(output_file, 'w') as f:
            f.write(prompt)
        print(f"Prompt saved to: {output_file}")
    
    def process(self, output_file="generated_prompt.txt"):
        """Main processing method."""
        print("Starting Excel formula analysis...")
        
        # Step 1: Load workbook
        self.load_workbook()
        
        # Step 2: Create header mapping
        self.create_header_mapping()
        
        # Step 3: Extract formulas
        self.extract_formulas()
        
        # Step 4: Extract data from non-formula columns
        self.extract_data_columns()
        
        # Step 5: Generate prompt
        prompt = self.generate_prompt()
        
        # Step 6: Save prompt
        self.save_prompt_to_file(prompt, output_file)
        
        return prompt


def main():
    parser = argparse.ArgumentParser(description='Generate LLM prompts for Excel formula conversion')
    parser.add_argument('excel_file', help='Path to Excel file with formulas')
    parser.add_argument('-o', '--output', default='generated_prompt.txt', 
                       help='Output file for generated prompt (default: generated_prompt.txt)')
    
    args = parser.parse_args()
    
    if not Path(args.excel_file).exists():
        print(f"Error: Excel file '{args.excel_file}' not found.")
        sys.exit(1)
    
    generator = ExcelFormulaPromptGenerator(args.excel_file)
    prompt = generator.process(args.output)
    
    print("\n" + "="*50)
    print("GENERATED PROMPT:")
    print("="*50)
    print(prompt)


if __name__ == "__main__":
    main() 